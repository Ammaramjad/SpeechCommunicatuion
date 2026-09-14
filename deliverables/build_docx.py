#!/usr/bin/env python3
"""Word companion of the boxed AMD–ITRI conversion (same substance as the PDF/MD)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "AMD-ITRI-MDA-TARI_MI300X-RAG_Boxed-Project-Review.docx"


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): hex_color, qn("w:val"): "clear"})
    tcPr.append(shd)


def set_run(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(11, 30, 58)
    return p


def para(doc, text, *, bold=False, size=11, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run(run, size=11)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run(r, size=10, bold=(i == 0 or j == 0))
            if i == 0:
                shade(cell, "0B1E3A")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("AMD – ITRI PROJECT REVIEW")
    set_run(r, size=12, bold=True, color="00A8D0")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MI300X Enterprise Document RAG Project Report")
    set_run(r, size=22, bold=True, color="0B1E3A")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Boxed / पेटी conversion onto the AMD–ITRI Project Review template")
    set_run(r, size=12, color="1A3A5C")

    para(doc, "Dr. Ammar Amjad  ·  National Yang Ming Chiao Tung University", size=12)
    para(doc, "AMD Instinct MI300X  ·  PyTorch 2.3 / ROCm 6.2  ·  Implemented RAG prototype", size=11)

    heading(doc, "Document control (attestation — not a legal signature)", 1)
    para(
        doc,
        "Prepared/generated 2026-09-11 as a document-control conversion. "
        "This block is not a handwritten, cryptographic, or legal personal signature.",
        size=10,
        color="666666",
    )
    add_table(
        doc,
        [
            ["Field", "Value"],
            ["Document title", "AMD–ITRI Project Review: MI300X Enterprise Document RAG (boxed / पेटी format)"],
            ["Template", "AMD–ITRI PROJECT REVIEW boxed slide template (navy + cyan rounded boxes)"],
            ["पेटी interpretation", "Official boxed layout of the Project Review deck, not a ZIP package"],
            ["एमडीए टारी interpretation", "AMD ITRI Project Review template (not an SEC MD&A filing)"],
            ["Version / date", "1.0  ·  2026-09-11"],
            ["Author (source)", "Dr. Ammar Amjad, NYCU"],
            ["Prepared from", "AMD-ITRI-Project-Review-TEMPLATE.pdf + AMD-ITRI-MI300X-Final-Comparison-Report-SOURCE.pdf"],
            ["Classification", "Project review · Technical report"],
            ["Signature", "None forged. Document-control attestation only."],
        ],
    )

    heading(doc, "1. Business problem and implemented solution", 1)
    heading(doc, "BUSINESS PROBLEM", 2)
    para(
        doc,
        "Enterprise knowledge is spread across PDF reports and Excel sheets. "
        "Finding the right information manually is slow, and answers often lack a traceable source.",
    )
    heading(doc, "IMPLEMENTED SOLUTION", 2)
    para(
        doc,
        "Retrieval-Augmented Generation (RAG): the system first retrieves the most relevant document chunks, "
        "then Qwen2.5-7B-Instruct generates a Traditional Chinese answer grounded in them.",
    )
    add_table(
        doc,
        [
            ["Step", "Content"],
            ["1 Data Preparation", "PDF + Excel → text extraction → 500-char chunks → BGE embeddings → vector index"],
            ["2 Retrieval", "User question → query embedding → dot-product similarity → Top-3 chunks"],
            ["3 Generation", "Retrieved context + sources → Qwen2.5-7B-Instruct → Traditional Chinese answer with filenames"],
        ],
    )

    heading(doc, "2. Project objective", 1)
    heading(doc, "PROBLEM ADDRESSED", 2)
    bullets(
        doc,
        [
            "Information distributed across PDF and Excel files",
            "Manual search is time-consuming",
            "Answers may lack document sources / need traceable document sources",
        ],
    )
    heading(doc, "IMPLEMENTED OBJECTIVE", 2)
    bullets(
        doc,
        [
            "Build a searchable document knowledge base / Chinese document knowledge index",
            "Retrieve relevant chunks for a question",
            "Generate source-grounded Traditional Chinese answers with Qwen",
        ],
    )

    heading(doc, "3. Knowledge index construction", 1)
    bullets(
        doc,
        [
            "PDF + Excel: enterprise source documents (PDF pages; Excel rows and sheets)",
            "Text extraction: pypdf (PDF); openpyxl (Excel rows)",
            "Chunking: clean text; 500-character chunks; 100-character overlap",
            "BGE embeddings: BAAI/bge-small-zh-v1.5; normalized, on MI300X",
            "Vector index: gama_embeddings.npy, gama_chunks.json, index_summary.json",
        ],
    )
    add_table(
        doc,
        [
            ["Index fact", "Value"],
            ["Document sections extracted", "25"],
            ["Searchable text chunks created", "17"],
            ["Embedding normalization", "Normalized for dot-product similarity"],
            ["Persistence", "Index persisted to disk for reuse across queries"],
        ],
    )

    heading(doc, "4. Implemented architecture", 1)
    para(doc, "OFFLINE · INDEX CONSTRUCTION", bold=True)
    para(doc, "Enterprise documents (PDF · Excel) → text extraction & chunking (pypdf · openpyxl · 500/100) → BGE embeddings (bge-small-zh-v1.5 · MI300X) → vector index (.npy + .json).")
    para(doc, "ONLINE · QUESTION ANSWERING", bold=True)
    para(doc, "User question (Traditional Chinese) → Top-3 retrieval (normalized · dot-product) → Qwen2.5-7B-Instruct (bfloat16 · device_map=auto) → grounded answer (Traditional Chinese + source filenames).")
    para(doc, "Runs on AMD Instinct MI300X · PyTorch 2.3 · ROCm 6.2. The system retrieves relevant document evidence before Qwen2.5-7B generates the answer.")
    heading(doc, "Additional facts from the comparison-source architecture diagram", 2)
    bullets(
        doc,
        [
            "Sentence Transformers; batch size 32; normalized vector embeddings",
            "Index metadata: source + page + sheet + row",
            "Similarity: score(q, d) = qᵀd on normalized vectors (cosine)",
            "Generation: max_new_tokens = 300",
            "Observed LLM generation time on the diagram: 2.55 seconds",
        ],
    )
    para(doc, "Source figure embedded in the PDF as assets/implemented-flow-source.png.")

    heading(doc, "5. Model roles", 1)
    heading(doc, "BAAI/bge-small-zh-v1.5 — Embedding and retrieval", 2)
    bullets(
        doc,
        [
            "Converts document chunks and user queries into semantic / comparable vectors",
            "Normalized vector search and Top-3 retrieval by dot-product similarity",
            "Chinese semantic embeddings",
        ],
    )
    heading(doc, "Qwen2.5-7B-Instruct — Answer generation", 2)
    bullets(
        doc,
        [
            "Uses retrieved context and system instructions",
            "Generates source-grounded Traditional Chinese answers and source filenames",
            "bfloat16 inference on MI300X",
        ],
    )
    para(doc, "No model training or fine-tuning is claimed in this project.", bold=True)

    heading(doc, "6. Verified results and comparison", 1)
    add_table(
        doc,
        [
            ["KPI", "Value"],
            ["Document sections extracted", "25"],
            ["Searchable chunks created", "17"],
            ["Observed Qwen generation time", "2.55 s (single-query observation, not a general GPU benchmark)"],
            ["Retrieved chunks per query", "Top-3"],
            ["Execution environment", "AMD Instinct MI300X · PyTorch 2.3 · ROCm 6.2"],
        ],
    )
    add_table(
        doc,
        [
            ["Model (comparison-source chart)", "Observed response time (seconds)"],
            ["Qwen2.5-7B", "2.55"],
            ["Llama 3.1 8B Instruct", "3.1"],
            ["Mistral 7B Instruct", "2.9"],
        ],
    )
    heading(doc, "OBSERVATIONS", 2)
    bullets(
        doc,
        [
            "Qwen2.5-7B-Instruct loaded and executed successfully on MI300X (bfloat16, device_map=auto).",
            "The answer used the retrieved documents and stated “insufficient information” instead of inventing a third source.",
        ],
    )

    heading(doc, "7. Proposed deployment model", 1)
    para(doc, "PROPOSED · NOT YET IMPLEMENTED", bold=True, color="00A8D0")
    bullets(
        doc,
        [
            "Controlled upload: authorized users upload PDF / Excel documents",
            "Versioning & indexing: track versions; re-embed and update the index",
            "Source-grounded Q&A: retrieve evidence, then generate the answer with Qwen",
            "Citations & feedback: show source filenames; collect user ratings",
        ],
    )
    para(doc, "Future-state interface (concept): Q: Ask a question about your enterprise documents…  A: Grounded answer · Sources: report.pdf, data.xlsx · Was this helpful? Yes / No")
    para(doc, "Proposed 12-box enterprise flow: see the markdown/PDF transcription and assets/proposed-deployment-source.png. This flow is proposed, not claimed as implemented.")

    heading(doc, "8. Request for renewed MI300X access", 1)
    heading(doc, "WHAT THE INITIAL ALLOCATION ACHIEVED", 2)
    para(doc, "The first MI300X allocation validated the technical feasibility of the full workflow on AMD hardware with PyTorch 2.3 / ROCm 6.2.")
    bullets(
        doc,
        [
            "Document indexing (PDF + Excel → 17 chunks)",
            "Normalized Top-3 semantic retrieval",
            "Qwen2.5-7B-Instruct inference (bfloat16)",
            "Source-grounded Traditional Chinese answers",
        ],
    )
    heading(doc, "PLANNED NEXT STEPS WITH RENEWED ACCESS", 2)
    bullets(
        doc,
        [
            "Expand official product and service documents",
            "Create a 50–100 question evaluation set",
            "Measure retrieval quality and inference performance reproducibly / systematically",
            "Build the proposed enterprise knowledge interface",
        ],
    )
    heading(doc, "EXPECTED DELIVERABLES (comparison source)", 2)
    bullets(
        doc,
        [
            "Reproducible MI300X benchmark report",
            "Validated source-grounded RAG results",
            "A proposed interface for enterprise knowledge support",
        ],
    )

    heading(doc, "Closing", 1)
    para(doc, "Thank you to AMD and ITRI for enabling this work.", bold=True)
    para(
        doc,
        "Prepared/generated conversion · 2026-09-11 · Document-control attestation only. "
        "No official or personal legal signature was applied.",
        size=10,
        color="666666",
    )

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
